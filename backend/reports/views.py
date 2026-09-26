from django.shortcuts import render
from django.contrib.auth.decorators import login_required
from django.http import HttpResponse
from django.db.models import Sum
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet
import io

from clients.models import Client
from pipeline.models import Deal

@login_required
def index(request):
    date_from = request.GET.get('date_from')
    date_to = request.GET.get('date_to')
    
    clients = Client.objects.all()
    deals = Deal.objects.all()
    
    if date_from:
        date_from_dt = datetime.strptime(date_from, '%Y-%m-%d')
        clients = clients.filter(created_at__gte=date_from_dt)
        deals = deals.filter(created_at__gte=date_from_dt)
        
    if date_to:
        date_to_dt = datetime.strptime(date_to, '%Y-%m-%d')
        clients = clients.filter(created_at__lte=date_to_dt)
        deals = deals.filter(created_at__lte=date_to_dt)
        
    clients_count = clients.count()
    deals_count = deals.count()
    
    won_deals = deals.filter(stage__name='Fechado Ganho')
    won_deals_count = won_deals.count()
    won_deals_value = won_deals.aggregate(total=Sum('value'))['total'] or 0
    
    lost_deals = deals.filter(stage__name='Perdido')
    lost_deals_count = lost_deals.count()
    
    context = {
        'date_from': date_from,
        'date_to': date_to,
        'clients_count': clients_count,
        'deals_count': deals_count,
        'won_deals_count': won_deals_count,
        'won_deals_value': won_deals_value,
        'lost_deals_count': lost_deals_count,
    }
    
    return render(request, 'reports/index.html', context)

@login_required
def export_clients_excel(request):
    wb = Workbook()
    ws = wb.active
    ws.title = "Clientes"
    
    headers = ['Nome', 'Email', 'Telefone', 'Empresa', 'Status', 'Data de Criação']
    ws.append(headers)
    
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4F46E5", end_color="4F46E5", fill_type="solid")
    
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        
    for client in Client.objects.all():
        ws.append([
            client.name,
            client.email,
            client.phone,
            client.company,
            client.get_status_display(),
            client.created_at.strftime('%d/%m/%Y %H:%M') if client.created_at else ''
        ])
        
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename=clientes.xlsx'
    wb.save(response)
    return response

@login_required
def export_deals_excel(request):
    wb = Workbook()
    ws = wb.active
    ws.title = "Oportunidades"
    
    headers = ['Título', 'Cliente', 'Valor', 'Etapa', 'Responsável', 'Data Prevista', 'Data de Criação']
    ws.append(headers)
    
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4F46E5", end_color="4F46E5", fill_type="solid")
    
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        
    for deal in Deal.objects.all():
        ws.append([
            deal.title,
            deal.client.name,
            float(deal.value),
            deal.stage.name if deal.stage else '',
            deal.owner.username if deal.owner else '',
            deal.expected_close_date.strftime('%d/%m/%Y') if deal.expected_close_date else '',
            deal.created_at.strftime('%d/%m/%Y %H:%M') if deal.created_at else ''
        ])
        
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename=oportunidades.xlsx'
    wb.save(response)
    return response

@login_required
def export_report_pdf(request):
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter)
    elements = []
    styles = getSampleStyleSheet()
    
    elements.append(Paragraph("Relatório de Vendas - CRM Vendas", styles['Title']))
    elements.append(Paragraph(f"Gerado em: {datetime.now().strftime('%d/%m/%Y %H:%M')}", styles['Normal']))
    elements.append(Spacer(1, 20))
    
    # Summary
    deals = Deal.objects.all()
    won_deals = deals.filter(stage__name='Fechado Ganho')
    won_value = won_deals.aggregate(total=Sum('value'))['total'] or 0
    
    summary_data = [
        ['Métrica', 'Valor'],
        ['Total de Clientes', str(Client.objects.count())],
        ['Total de Oportunidades', str(deals.count())],
        ['Oportunidades Ganhas', str(won_deals.count())],
        ['Valor Total Ganho', f"R$ {won_value:,.2f}".replace(',', 'X').replace('.', ',').replace('X', '.')],
    ]
    
    t_summary = Table(summary_data, colWidths=[200, 200])
    t_summary.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#4F46E5')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 12),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
        ('BACKGROUND', (0, 1), (-1, -1), colors.HexColor('#f3f4f6')),
        ('GRID', (0, 0), (-1, -1), 1, colors.black),
    ]))
    
    elements.append(Paragraph("Resumo", styles['Heading2']))
    elements.append(t_summary)
    elements.append(Spacer(1, 20))
    
    # Deals Table
    deals_data = [['Título', 'Cliente', 'Valor', 'Etapa']]
    for deal in deals.order_by('-created_at')[:20]:
        deals_data.append([
            deal.title[:30],
            deal.client.name[:20],
            f"R$ {deal.value:,.2f}".replace(',', 'X').replace('.', ',').replace('X', '.'),
            deal.stage.name if deal.stage else ''
        ])
        
    t_deals = Table(deals_data, colWidths=[150, 150, 100, 100])
    t_deals.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#4F46E5')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
        ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
        ('ALIGN', (2, 1), (2, -1), 'RIGHT'),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 10),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 10),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
    ]))
    
    elements.append(Paragraph("Últimas Oportunidades", styles['Heading2']))
    elements.append(t_deals)
    
    doc.build(elements)
    buffer.seek(0)
    
    response = HttpResponse(buffer, content_type='application/pdf')
    response['Content-Disposition'] = 'attachment; filename="relatorio_vendas.pdf"'
    return response
